import logging
from datetime import timedelta
from io import BytesIO
from pathlib import Path
from urllib.parse import quote

import requests
from celery import shared_task
from django.conf import settings
from django.core.files.base import ContentFile
from django.db.models import Q
from django.utils import timezone
from docx import Document
from docx.shared import Pt

from apps.employees.models import EmployeeProfile

from .models import AttendanceTelegramDelivery, EmployeeTelegramAccount, WeeklyReport
from .telegram import UTC_PLUS_5, _safe_delay, employee_name, resolve_group_topic


logger = logging.getLogger(__name__)
TEMPLATE_PATH = Path(__file__).resolve().parent / 'templates' / 'weekly_report_template.docx'
PROMPTS = (
    ('Что проделано за неделю?', 'work_done'),
    ('Замечания', 'remarks'),
    ('Планы на следующую неделю', 'next_week_plans'),
    ('Что внедрять? Какие идеи есть для улучшения?', 'improvement_ideas'),
    ('Какие сложности есть?', 'difficulties'),
    ('Что нужно тебе может?', 'needs'),
    ('В пару с кем? От кого чего ждёт? (каталоги, обучение, макет, …?)', 'waiting_for'),
    ('Для информации, исключения', 'information'),
)


def current_week_bounds(day=None):
    day = day or timezone.now().astimezone(UTC_PLUS_5).date()
    start = day - timedelta(days=day.weekday())
    return start, start + timedelta(days=5)


def _style_run(run, *, bold=False):
    run.bold = bold
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)


def _replace_header(paragraph, report):
    paragraph.clear()
    name = employee_name(report.employee)
    office = str(report.office) if report.office_id else 'Без офиса'
    label = paragraph.add_run('ФИО: ')
    _style_run(label, bold=True)
    _style_run(paragraph.add_run(name))
    _style_run(paragraph.add_run('                 офис: '), bold=True)
    _style_run(paragraph.add_run(office))


def _replace_date(paragraph, report):
    paragraph.clear()
    _style_run(paragraph.add_run('дата: '), bold=True)
    _style_run(paragraph.add_run(report.period_end.strftime('%d.%m.%Y')))


def render_weekly_report(report):
    document = Document(str(TEMPLATE_PATH))
    _replace_header(document.paragraphs[0], report)
    _replace_date(document.paragraphs[1], report)
    cells = [cell for row in document.tables[0].rows for cell in row.cells]
    for cell, (prompt, field_name) in zip(cells, PROMPTS):
        answer = str(getattr(report, field_name) or '').strip() or '—'
        cell._tc.clear_content()
        paragraph = cell.add_paragraph()
        _style_run(paragraph.add_run(prompt), bold=True)
        paragraph.add_run('\n\n')
        _style_run(paragraph.add_run(answer))
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def save_generated_report(report):
    content = render_weekly_report(report)
    filename = f'weekly-report-{report.employee_id}-{report.period_end:%Y-%m-%d}.docx'
    if not report.generated_file:
        report.generated_file.save(filename, ContentFile(content), save=False)
        report.save(update_fields=['generated_file', 'updated_at'])
    return content, filename


def archive_report_to_disk(report, content, filename):
    endpoint = str(getattr(settings, 'DISK_REPORTS_API_URL', '') or '').strip()
    token = str(getattr(settings, 'DISK_PROVISION_SERVICE_TOKEN', '') or '').strip()
    if not endpoint or not token:
        raise RuntimeError('DiskSL report archive is not configured')
    response = requests.post(
        endpoint,
        data=content,
        headers={
            'Authorization': f'Bearer {token}',
            'Content-Type': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            'X-File-Name': quote(filename),
            'X-Employee-Name': quote(employee_name(report.employee)),
            'X-Report-Period': f'{report.period_start:%Y-%m-%d}_{report.period_end:%Y-%m-%d}',
            'X-Report-ID': str(report.pk),
        },
        timeout=(5, 120),
    )
    response.raise_for_status()
    payload = response.json()
    report.disk_path = str(payload.get('path') or '')[:500]
    report.disk_archived_at = timezone.now()
    report.save(update_fields=['disk_path', 'disk_archived_at', 'updated_at'])


def send_report_to_telegram(report, content, filename):
    topic = resolve_group_topic(AttendanceTelegramDelivery.EVENT_REPORT_SUBMITTED)
    if not topic:
        raise RuntimeError('telegram_topic_not_configured:reports')
    token = settings.ATTENDANCE_TELEGRAM_BOT_TOKEN
    base = settings.ATTENDANCE_TELEGRAM_DOCUMENT_API_BASE.rstrip('/')
    caption = (
        '#еженедельный_отчёт\n'
        f'{employee_name(report.employee)}\n'
        f'{report.period_start:%d.%m.%Y}–{report.period_end:%d.%m.%Y} · UTC+5'
    )
    response = requests.post(
        f'{base}/bot{token}/sendDocument',
        data={
            'chat_id': topic.chat_id,
            'message_thread_id': topic.message_thread_id,
            'caption': caption,
        },
        files={
            'document': (
                filename,
                content,
                'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            ),
        },
        timeout=(5, 120),
    )
    data = response.json()
    if response.status_code != 200 or data.get('ok') is not True:
        raise RuntimeError(f'Telegram HTTP {response.status_code}')
    report.telegram_sent_at = timezone.now()
    report.telegram_error = ''
    report.save(update_fields=['telegram_sent_at', 'telegram_error', 'updated_at'])


@shared_task(name='attendance.process_weekly_report')
def process_weekly_report(report_id):
    report = WeeklyReport.objects.select_related('employee', 'office').get(pk=report_id)
    content, filename = save_generated_report(report)
    errors = []
    if not report.disk_archived_at:
        try:
            archive_report_to_disk(report, content, filename)
        except Exception as exc:
            logger.warning('Weekly report %s was not archived to DiskSL: %s', report.pk, type(exc).__name__)
            errors.append(f'DiskSL: {type(exc).__name__}')
    if not report.telegram_sent_at:
        try:
            send_report_to_telegram(report, content, filename)
        except Exception as exc:
            logger.warning('Weekly report %s was not sent to Telegram: %s', report.pk, type(exc).__name__)
            errors.append(f'Telegram: {type(exc).__name__}')
    if errors:
        report.telegram_error = '; '.join(errors)[:255]
        report.save(update_fields=['telegram_error', 'updated_at'])
    return {'report_id': report.pk, 'errors': errors}


@shared_task(name='attendance.retry_weekly_report_delivery')
def retry_weekly_report_delivery():
    ids = list(
        WeeklyReport.objects.filter(expires_at__gt=timezone.now())
        .filter(Q(telegram_sent_at__isnull=True) | Q(disk_archived_at__isnull=True))
        .order_by('submitted_at')
        .values_list('pk', flat=True)[:20]
    )
    for report_id in ids:
        process_weekly_report.delay(report_id)
    return {'queued': len(ids)}


@shared_task(name='attendance.cleanup_weekly_report_archives')
def cleanup_weekly_report_archives():
    reports = list(WeeklyReport.objects.filter(expires_at__lt=timezone.now()).exclude(generated_file='')[:100])
    deleted = 0
    endpoint = str(getattr(settings, 'DISK_REPORTS_API_URL', '') or '').strip()
    token = str(getattr(settings, 'DISK_PROVISION_SERVICE_TOKEN', '') or '').strip()
    for report in reports:
        if report.disk_path and endpoint and token:
            try:
                requests.delete(
                    endpoint,
                    params={'path': report.disk_path},
                    headers={'Authorization': f'Bearer {token}'},
                    timeout=(5, 30),
                ).raise_for_status()
            except requests.RequestException:
                continue
        report.generated_file.delete(save=False)
        report.generated_file = ''
        report.disk_path = ''
        report.save(update_fields=['generated_file', 'disk_path', 'updated_at'])
        deleted += 1
    return {'deleted': deleted}


@shared_task(name='attendance.remind_weekly_reports')
def remind_weekly_reports():
    week_start, _ = current_week_bounds()
    if timezone.now().astimezone(UTC_PLUS_5).date().weekday() != 5:
        return {'queued': 0, 'reason': 'not_saturday'}
    profiles = EmployeeProfile.objects.filter(
        is_active=True,
        work_status='working',
        user__is_active=True,
    ).exclude(user__is_superuser=True).exclude(user__role='admin').select_related('user', 'company', 'office')
    queued = 0
    for profile in profiles:
        if WeeklyReport.objects.filter(employee=profile.user, company=profile.company, period_start=week_start).exists():
            continue
        account = EmployeeTelegramAccount.objects.filter(employee=profile.user, is_active=True).first()
        if not account:
            continue
        delivery, created = AttendanceTelegramDelivery.objects.get_or_create(
            event_key=f'personal-weekly-report:{profile.user_id}:{week_start:%Y-%m-%d}',
            defaults={
                'event_type': AttendanceTelegramDelivery.EVENT_REPORT_REMINDER,
                'company': profile.company,
                'office': profile.office,
                'employee': profile.user,
                'target_chat_id': account.chat_id,
                'message': (
                    '#еженедельный_отчёт\nСегодня суббота. Заполните недельный отчёт '
                    'в ManagerSL → Отчёты. После сохранения документ отправится в тему «Отчёты».'
                ),
            },
        )
        if created:
            _safe_delay(delivery.pk)
            queued += 1
    return {'queued': queued}
