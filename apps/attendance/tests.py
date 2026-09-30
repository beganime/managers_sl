from datetime import date, datetime, time, timedelta
from io import BytesIO
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.test import TestCase, override_settings
from django.utils import timezone
from docx import Document

from apps.employees.models import EmployeeProfile, EmployeeRole
from apps.erp_notifications.models import NotificationTemplate
from apps.erp_notifications.tasks import auto_close_workdays, send_attendance_reminders
from apps.organizations.models import Company

from .models import (
    AttendanceReminder,
    AttendanceTelegramDelivery,
    AttendanceTelegramTopic,
    EmployeeTelegramAccount,
    WorkDay,
    WeeklyReport,
)
from .services import auto_start_workday_for_login, is_scheduled_workday, record_after_hours_activity, workday_close_at
from .telegram import (
    create_employee_link,
    daily_summary_messages,
    queue_admin_message,
    register_personal_reminder,
    retry_attendance_telegram_deliveries,
    send_daily_attendance_summary,
    send_attendance_telegram_delivery,
    send_weekly_attendance_summary,
    weekly_summary_messages,
    workday_message,
)
from .weekly_reports import render_weekly_report


@override_settings(ATTENDANCE_WORKDAYS=(0, 1, 2, 3, 4, 5), ATTENDANCE_AUTO_CLOSE_HOUR=20)
class AttendanceScheduleTests(TestCase):
    def test_sunday_is_not_a_workday(self):
        self.assertFalse(is_scheduled_workday(date(2026, 10, 4)))

    def test_automatic_close_is_at_twenty_hundred(self):
        self.assertEqual(workday_close_at(date(2026, 9, 28)).hour, 20)


class WeeklyReportDocumentTests(TestCase):
    def test_reference_template_is_filled_with_employee_answers(self):
        company = Company.objects.create(name='Students Life', country='Туркменистан')
        user = get_user_model().objects.create_user(
            email='weekly@example.com',
            password='test-password',
            first_name='Анна',
            last_name='Иванова',
        )
        report = WeeklyReport(
            company=company,
            employee=user,
            period_start=date(2026, 9, 28),
            period_end=date(2026, 10, 3),
            work_done='Проверила документы клиентов.',
            next_week_plans='Подготовить новые заявки.',
            expires_at=timezone.now() + timedelta(days=30),
        )

        rendered = Document(BytesIO(render_weekly_report(report)))

        self.assertIn('Анна', rendered.paragraphs[0].text)
        self.assertIn('03.10.2026', rendered.paragraphs[1].text)
        self.assertIn('Проверила документы клиентов.', rendered.tables[0].cell(0, 0).text)
        self.assertIn('Подготовить новые заявки.', rendered.tables[0].cell(1, 0).text)


@override_settings(ATTENDANCE_WORKDAYS=(0, 1, 2, 3, 4, 5, 6), ATTENDANCE_ACTIVITY_PROTECTION_HOUR=17)
class WorkdayClosingPolicyTests(TestCase):
    def setUp(self):
        self.company = Company.objects.create(name='Students Life', country='Туркменистан')
        self.role = EmployeeRole.objects.create(code='manager', name='Менеджер', role_type='manager')

    def create_employee(self, email):
        user = get_user_model().objects.create_user(email=email, password='test-password')
        EmployeeProfile.objects.create(user=user, company=self.company, role=self.role)
        return user

    def set_activity(self, user, hour, minute=0):
        local_value = timezone.make_aware(
            datetime.combine(timezone.localdate(), time(hour, minute)),
            timezone.get_current_timezone(),
        )
        type(user).objects.filter(pk=user.pk).update(last_activity=local_value)

    @patch('apps.erp_notifications.tasks.create_notification')
    def test_employee_without_start_is_recorded_as_missed(self, _create_notification):
        user = self.create_employee('absent@example.com')

        result = auto_close_workdays()

        workday = WorkDay.objects.get(employee=user, date=timezone.localdate())
        self.assertEqual(workday.status, WorkDay.STATUS_MISSED)
        self.assertEqual(result['missed'], 1)

    @patch('apps.erp_notifications.tasks.create_notification')
    def test_inactive_started_day_is_closed_at_20(self, _create_notification):
        user = self.create_employee('inactive@example.com')
        self.set_activity(user, 16, 59)
        workday = WorkDay.objects.create(
            company=self.company,
            employee=user,
            date=timezone.localdate(),
            status=WorkDay.STATUS_STARTED,
            started_at=timezone.now(),
        )

        result = auto_close_workdays()

        workday.refresh_from_db()
        self.assertEqual(workday.status, WorkDay.STATUS_AUTO_CLOSED)
        self.assertEqual(result['closed'], 1)

    @patch('apps.erp_notifications.tasks.create_notification')
    def test_evening_activity_is_closed_by_automatic_run_at_20(self, _create_notification):
        user = self.create_employee('active@example.com')
        self.set_activity(user, 17, 30)
        workday = WorkDay.objects.create(
            company=self.company,
            employee=user,
            date=timezone.localdate(),
            status=WorkDay.STATUS_STARTED,
            started_at=timezone.now(),
        )

        result = auto_close_workdays()

        workday.refresh_from_db()
        self.assertEqual(workday.status, WorkDay.STATUS_AUTO_CLOSED)
        self.assertFalse(workday.requires_manual_close)
        self.assertEqual(result['closed'], 1)

    def test_activity_after_20_is_recorded_without_reopening_day(self):
        user = self.create_employee('late@example.com')
        local_value = timezone.make_aware(
            datetime.combine(timezone.localdate(), time(20, 25)),
            timezone.get_current_timezone(),
        )
        workday = WorkDay.objects.create(
            company=self.company,
            employee=user,
            date=timezone.localdate(),
            status=WorkDay.STATUS_AUTO_CLOSED,
            started_at=local_value - timedelta(hours=9),
            closed_at=local_value - timedelta(minutes=25),
        )

        recorded, created = record_after_hours_activity(user, local_value)

        self.assertEqual(recorded.pk, workday.pk)
        self.assertTrue(created)
        recorded.refresh_from_db()
        self.assertEqual(recorded.status, WorkDay.STATUS_AUTO_CLOSED)
        self.assertEqual(recorded.custom_data['after_hours_activity_count'], 1)
        self.assertIn('after_hours_first_seen_at', recorded.custom_data)

        record_after_hours_activity(user, local_value + timedelta(seconds=20))
        recorded.refresh_from_db()
        self.assertEqual(recorded.custom_data['after_hours_activity_count'], 1)

    @patch('apps.erp_notifications.tasks.ensure_default_attendance_reminders')
    @patch('apps.erp_notifications.tasks.create_notification')
    def test_start_push_targets_every_active_employee(self, create_notification, _ensure_defaults):
        first = self.create_employee('first@example.com')
        second = self.create_employee('second@example.com')
        now = timezone.localtime()
        AttendanceReminder.objects.create(
            company=self.company,
            reminder_type=AttendanceReminder.REMINDER_START,
            scheduled_time=now.time().replace(second=0, microsecond=0),
            weekdays=[now.weekday()],
            message='Начните рабочий день.',
        )

        created = send_attendance_reminders(AttendanceReminder.REMINDER_START)

        self.assertEqual(created, 2)
        self.assertEqual({call.args[0] for call in create_notification.call_args_list}, {first, second})
        self.assertTrue(
            all(call.kwargs['channel'] == NotificationTemplate.CHANNEL_PUSH for call in create_notification.call_args_list)
        )

    @patch('apps.erp_notifications.tasks.ensure_default_attendance_reminders')
    @patch('apps.erp_notifications.tasks.create_notification')
    def test_administrator_gets_no_attendance_reminder_or_automatic_absence(self, create_notification, _ensure_defaults):
        administrator = get_user_model().objects.create_user(
            email='attendance-admin@example.com', password='test-password', role='admin', is_staff=True,
        )
        EmployeeProfile.objects.create(user=administrator, company=self.company, role=self.role)
        now = timezone.localtime()
        AttendanceReminder.objects.create(
            company=self.company,
            reminder_type=AttendanceReminder.REMINDER_START,
            scheduled_time=now.time().replace(second=0, microsecond=0),
            weekdays=[now.weekday()],
            message='Начните рабочий день.',
        )

        reminders = send_attendance_reminders(AttendanceReminder.REMINDER_START)
        close_result = auto_close_workdays()

        self.assertEqual(reminders, 0)
        self.assertEqual(close_result['missed'], 0)
        create_notification.assert_not_called()
        self.assertFalse(WorkDay.objects.filter(employee=administrator).exists())


@override_settings(
    ATTENDANCE_WORKDAYS=(0, 1, 2, 3, 4, 5, 6),
    ATTENDANCE_TELEGRAM_ENABLED=True,
    ATTENDANCE_TELEGRAM_BOT_TOKEN='test-token',
    ATTENDANCE_TELEGRAM_CHAT_ID='-1001',
    ATTENDANCE_TELEGRAM_BOT_USERNAME='manager_sl_test_bot',
    ATTENDANCE_TELEGRAM_WEBHOOK_SECRET='webhook-test-secret',
)
class AttendanceTelegramTests(TestCase):
    def setUp(self):
        self.company = Company.objects.create(name='Students Life', country='Туркменистан')
        self.role = EmployeeRole.objects.create(code='manager', name='Менеджер', role_type='manager')
        self.user = get_user_model().objects.create_user(
            email='manager@example.com',
            password='test-password',
            first_name='Анна',
            last_name='Иванова',
        )
        EmployeeProfile.objects.create(user=self.user, company=self.company, role=self.role)

    @patch('apps.attendance.services.is_after_workday_close', return_value=False)
    def test_login_starts_workday_and_arrival_is_registered_once(self, _after_close):
        first, first_started = auto_start_workday_for_login(self.user)
        second, second_started = auto_start_workday_for_login(self.user)

        self.assertTrue(first_started)
        self.assertFalse(second_started)
        self.assertEqual(first.pk, second.pk)
        self.assertEqual(first.status, WorkDay.STATUS_STARTED)
        self.assertEqual(
            AttendanceTelegramDelivery.objects.filter(event_type=AttendanceTelegramDelivery.EVENT_ARRIVAL).count(),
            1,
        )

    def test_administrator_is_excluded_from_workday_and_team_summaries(self):
        administrator = get_user_model().objects.create_user(
            email='admin@example.com', password='test-password', first_name='Администратор', role='admin', is_staff=True,
        )
        EmployeeProfile.objects.create(user=administrator, company=self.company, role=self.role)

        workday, started = auto_start_workday_for_login(administrator)
        daily_message = '\n'.join(daily_summary_messages(self.company, timezone.localdate()))
        weekly_message = '\n'.join(
            weekly_summary_messages(
                self.company,
                timezone.localdate() - timedelta(days=timezone.localdate().weekday()),
                timezone.localdate(),
            )
        )

        self.assertIsNone(workday)
        self.assertFalse(started)
        self.assertFalse(WorkDay.objects.filter(employee=administrator).exists())
        self.assertNotIn('Администратор', daily_message)
        self.assertNotIn('Администратор', weekly_message)

    @patch('apps.attendance.services.is_after_workday_close', return_value=False)
    def test_staff_manager_still_has_to_track_workday(self, _after_close):
        staff_manager = get_user_model().objects.create_user(
            email='staff-manager@example.com', password='test-password', first_name='Старший менеджер', is_staff=True,
        )
        EmployeeProfile.objects.create(user=staff_manager, company=self.company, role=self.role)

        workday, started = auto_start_workday_for_login(staff_manager)
        daily_message = '\n'.join(daily_summary_messages(self.company, timezone.localdate()))

        self.assertTrue(started)
        self.assertEqual(workday.status, WorkDay.STATUS_STARTED)
        self.assertIn('Старший менеджер', daily_message)

    def test_manual_and_automatic_close_have_distinct_events(self):
        workday, _ = auto_start_workday_for_login(self.user)
        workday.close(comment='Завершил работу.')

        self.assertTrue(
            AttendanceTelegramDelivery.objects.filter(
                workday=workday,
                event_type=AttendanceTelegramDelivery.EVENT_DEPARTURE,
            ).exists()
        )

    @patch('apps.erp_notifications.tasks.create_notification')
    def test_missing_employee_registers_absence_event(self, _create_notification):
        result = auto_close_workdays()

        self.assertEqual(result['missed'], 1)
        self.assertTrue(
            AttendanceTelegramDelivery.objects.filter(
                employee=self.user,
                event_type=AttendanceTelegramDelivery.EVENT_MISSED,
            ).exists()
        )

    def test_weekly_summary_is_grouped_and_created_only_once(self):
        today = timezone.localdate()
        period_start = today - timedelta(days=today.weekday())
        message = '\n'.join(weekly_summary_messages(self.company, period_start, today))
        self.assertIn('Анна Иванова', message)
        self.assertIn('#недельный_отчёт', message)

        send_weekly_attendance_summary()
        send_weekly_attendance_summary()
        self.assertEqual(
            AttendanceTelegramDelivery.objects.filter(event_type=AttendanceTelegramDelivery.EVENT_WEEKLY).count(),
            1,
        )

    def test_daily_summary_lists_started_and_not_started_in_utc_plus_five(self):
        second = get_user_model().objects.create_user(
            email='second@example.com', password='test-password', first_name='Борис', last_name='Петров',
        )
        EmployeeProfile.objects.create(user=second, company=self.company, role=self.role)
        started_at = timezone.make_aware(datetime.combine(timezone.localdate(), time(9, 15)))
        WorkDay.objects.create(
            company=self.company,
            employee=self.user,
            date=timezone.localdate(),
            status=WorkDay.STATUS_STARTED,
            started_at=started_at,
        )

        message = '\n'.join(daily_summary_messages(self.company, timezone.localdate()))

        self.assertIn('Вышли на работу: 1', message)
        self.assertIn('Анна Иванова — 09:15 (UTC+5)', message)
        self.assertIn('Не вышли: 1', message)
        self.assertIn('Борис Петров', message)

        send_daily_attendance_summary()
        send_daily_attendance_summary()
        self.assertEqual(
            AttendanceTelegramDelivery.objects.filter(event_type=AttendanceTelegramDelivery.EVENT_DAILY).count(),
            1,
        )

    @override_settings(ATTENDANCE_WORKDAYS=(0, 1, 2, 3, 4, 5))
    @patch('apps.attendance.telegram.utc5_today', return_value=date(2026, 10, 4))
    def test_daily_summary_is_not_created_on_sunday(self, _today):
        result = send_daily_attendance_summary()

        self.assertEqual(result['reason'], 'non_working_day')
        self.assertFalse(AttendanceTelegramDelivery.objects.filter(event_type=AttendanceTelegramDelivery.EVENT_DAILY).exists())

    def test_workday_event_time_has_explicit_utc_plus_five_label(self):
        workday = WorkDay.objects.create(
            company=self.company,
            employee=self.user,
            date=timezone.localdate(),
            status=WorkDay.STATUS_STARTED,
            started_at=timezone.now(),
        )
        self.assertIn('(UTC+5)', workday_message(workday, AttendanceTelegramDelivery.EVENT_ARRIVAL))

    @patch('apps.attendance.telegram._enqueue')
    def test_report_reminder_and_admin_message_are_queued_with_distinct_events(self, enqueue):
        EmployeeTelegramAccount.objects.create(
            employee=self.user,
            telegram_user_id=7101,
            chat_id=7102,
        )
        reminder = register_personal_reminder(self.user, 'daily_report', timezone.localdate())
        administrator = get_user_model().objects.create_user(
            email='boss@example.com', password='test-password', first_name='Башлык', role='admin', is_staff=True,
        )
        EmployeeProfile.objects.create(user=administrator, company=self.company, role=self.role)
        admin_delivery = queue_admin_message(administrator, 'Кто сейчас свободен?')

        self.assertEqual(reminder.event_type, AttendanceTelegramDelivery.EVENT_REPORT_REMINDER)
        self.assertIn('17:30', reminder.message)
        self.assertEqual(admin_delivery.event_type, AttendanceTelegramDelivery.EVENT_ADMIN_MESSAGE)
        self.assertIn('#сообщение_руководителя', admin_delivery.message)
        self.assertIn('Кто сейчас свободен?', admin_delivery.message)
        self.assertEqual(enqueue.call_count, 2)

    def test_employee_links_telegram_with_one_time_start_code(self):
        link = create_employee_link(self.user)
        code = link.split('start=', 1)[1]
        response = self.client.post(
            '/api/integrations/telegram/attendance/webhook/',
            data={
                'message': {
                    'text': f'/start {code}',
                    'chat': {'id': 7002, 'type': 'private'},
                    'from': {'id': 7001, 'username': 'anna_manager', 'first_name': 'Анна'},
                },
            },
            content_type='application/json',
            HTTP_X_TELEGRAM_BOT_API_SECRET_TOKEN='webhook-test-secret',
        )

        self.assertEqual(response.status_code, 200)
        account = EmployeeTelegramAccount.objects.get(employee=self.user)
        self.assertEqual(account.telegram_user_id, 7001)
        self.assertEqual(account.chat_id, 7002)

        repeated = self.client.post(
            '/api/integrations/telegram/attendance/webhook/',
            data={
                'message': {
                    'text': f'/start {code}',
                    'chat': {'id': 7002, 'type': 'private'},
                    'from': {'id': 7001},
                },
            },
            content_type='application/json',
            HTTP_X_TELEGRAM_BOT_API_SECRET_TOKEN='webhook-test-secret',
        )
        self.assertEqual(repeated.status_code, 200)
        self.assertIn('недействительна', repeated.json()['text'])

    def test_forum_topic_is_registered_from_inside_the_topic(self):
        response = self.client.post(
            '/api/integrations/telegram/attendance/webhook/',
            data={
                'message': {
                    'text': '/topic reports',
                    'message_thread_id': 42,
                    'chat': {'id': -1001, 'type': 'supergroup'},
                    'from': {'id': 7001},
                },
            },
            content_type='application/json',
            HTTP_X_TELEGRAM_BOT_API_SECRET_TOKEN='webhook-test-secret',
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['message_thread_id'], 42)
        self.assertTrue(
            AttendanceTelegramTopic.objects.filter(
                chat_id=-1001,
                topic_type=AttendanceTelegramTopic.TOPIC_REPORTS,
                message_thread_id=42,
            ).exists()
        )

    def test_forum_topic_adopts_new_supergroup_id_after_forum_conversion(self):
        response = self.client.post(
            '/api/integrations/telegram/attendance/webhook/',
            data={
                'message': {
                    'text': '/topic daily',
                    'message_thread_id': 84,
                    'chat': {'id': -1005484558620, 'type': 'supergroup'},
                    'from': {'id': 7001},
                },
            },
            content_type='application/json',
            HTTP_X_TELEGRAM_BOT_API_SECRET_TOKEN='webhook-test-secret',
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['message_thread_id'], 84)
        self.assertTrue(
            AttendanceTelegramTopic.objects.filter(
                chat_id=-1005484558620,
                topic_type=AttendanceTelegramTopic.TOPIC_DAILY,
            ).exists()
        )

    def test_forum_topic_rejects_another_group_after_initial_binding(self):
        AttendanceTelegramTopic.objects.create(
            chat_id=-1005484558620,
            topic_type=AttendanceTelegramTopic.TOPIC_DAILY,
            message_thread_id=84,
        )
        response = self.client.post(
            '/api/integrations/telegram/attendance/webhook/',
            data={
                'message': {
                    'text': '/topic reports',
                    'message_thread_id': 91,
                    'chat': {'id': -100999, 'type': 'supergroup'},
                    'from': {'id': 7001},
                },
            },
            content_type='application/json',
            HTTP_X_TELEGRAM_BOT_API_SECRET_TOKEN='webhook-test-secret',
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json(), {'ok': True})
        self.assertFalse(AttendanceTelegramTopic.objects.filter(chat_id=-100999).exists())

    @patch('apps.attendance.telegram.send_telegram_message')
    def test_group_delivery_never_falls_back_to_general(self, send_message):
        delivery = AttendanceTelegramDelivery.objects.create(
            event_key='daily:no-topic',
            event_type=AttendanceTelegramDelivery.EVENT_DAILY,
            company=self.company,
            message='Итоги',
        )

        result = send_attendance_telegram_delivery(delivery.pk)

        self.assertFalse(result['sent'])
        send_message.assert_not_called()
        delivery.refresh_from_db()
        self.assertEqual(delivery.last_error, 'telegram_topic_not_configured:daily')

    @patch('apps.attendance.telegram.send_telegram_message', return_value=(True, ''))
    def test_group_delivery_uses_configured_forum_topic(self, send_message):
        AttendanceTelegramTopic.objects.create(
            chat_id=-1001,
            topic_type=AttendanceTelegramTopic.TOPIC_DAILY,
            message_thread_id=77,
        )
        delivery = AttendanceTelegramDelivery.objects.create(
            event_key='daily:with-topic',
            event_type=AttendanceTelegramDelivery.EVENT_DAILY,
            company=self.company,
            message='Итоги',
        )

        result = send_attendance_telegram_delivery(delivery.pk)

        self.assertTrue(result['sent'])
        send_message.assert_called_once_with('Итоги', -1001, 77)
        delivery.refresh_from_db()
        self.assertEqual(delivery.target_message_thread_id, 77)

    @patch('apps.attendance.telegram._safe_delay')
    def test_old_unrouted_delivery_is_not_replayed_when_topic_is_bound(self, enqueue):
        delivery = AttendanceTelegramDelivery.objects.create(
            event_key='daily:old-no-topic',
            event_type=AttendanceTelegramDelivery.EVENT_DAILY,
            company=self.company,
            message='Старая сводка',
            status=AttendanceTelegramDelivery.STATUS_FAILED,
            last_error='telegram_topic_not_configured:daily',
        )
        AttendanceTelegramDelivery.objects.filter(pk=delivery.pk).update(
            created_at=timezone.now() - timedelta(days=1),
        )

        result = retry_attendance_telegram_deliveries()

        self.assertEqual(result['queued'], 0)
        enqueue.assert_not_called()
